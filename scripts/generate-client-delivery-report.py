from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "CLIENT_PRODUCT_DELIVERY_REPORT.md"
OUTPUT = ROOT / "docs" / "Malabar_Coast_Complete_Product_Delivery_Report.docx"
LOGO = ROOT / "public" / "malabar.png"
HERO = ROOT / "public" / "malabar-restaurant-hero-v2.jpg"

INK = "10202A"
GREEN = "1E553D"
GOLD = "BC8D42"
CREAM = "F4EFE6"
MIST = "E7EEE9"
WHITE = "FFFFFF"
GREY = "5F6B70"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text_color(cell, color: str) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(color)


def add_page_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def add_hyperlink(paragraph, label: str, url: str) -> None:
    part = paragraph.part
    relationship = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), GREEN)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    text = OxmlElement("w:t")
    text.text = label
    run.extend([props, text])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\)|https?://[^\s)]+)")


def add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Aptos Mono"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(GREEN)
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        else:
            clean = token.rstrip(".,;")
            add_hyperlink(paragraph, clean, clean)
            if len(clean) < len(token):
                paragraph.add_run(token[len(clean) :])
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color, before, after in (
        ("Title", 34, WHITE, 0, 14),
        ("Subtitle", 15, WHITE, 0, 10),
        ("Heading 1", 23, GREEN, 18, 8),
        ("Heading 2", 15, GREEN, 13, 5),
        ("Heading 3", 11, GOLD, 9, 3),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    styles["List Bullet"].font.name = "Aptos"
    styles["List Bullet"].font.size = Pt(9.5)
    styles["List Number"].font.name = "Aptos"
    styles["List Number"].font.size = Pt(9.5)


def configure_page(section) -> None:
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.24)
    section.footer_distance = Inches(0.25)


def add_header_footer(section) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = "MALABAR COAST  ·  PRODUCT DELIVERY & OPERATIONS"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(GREEN)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("FINAL CLIENT HANDOVER  ·  29 SEPTEMBER 2026  ·  PAGE ")
    run.font.name = "Aptos"
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor.from_string(GREY)
    add_page_field(paragraph, "PAGE")


def add_cover(document: Document) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.95)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cell, INK)
    set_cell_margins(cell, top=540, start=460, bottom=540, end=460)

    if LOGO.exists():
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(2.7))
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("FINAL CONSOLIDATED PRODUCT DELIVERY,\nOPERATIONS & HANDOVER REPORT")
    run.font.name = "Aptos Display"
    run.font.size = Pt(29)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(WHITE)
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(16)

    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Website · Administration · CMS · Commerce · Infrastructure")
    run.font.name = "Aptos"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string("D9E7DF")

    if HERO.exists():
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(26)
        p.add_run().add_picture(str(HERO), width=Inches(5.9))

    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    run = p.add_run("Prepared for client handover\n29 September 2026  ·  Version 2.0\nwww.malabarcoast.co.uk")
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(WHITE)

    document.add_page_break()


def add_toc(document: Document) -> None:
    document.add_heading("Contents", level=1)
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Open this file in Microsoft Word and choose References → Update Table to refresh page numbers.")
    run.italic = True
    run.font.color.rgb = RGBColor.from_string(GREY)
    toc = document.add_paragraph()
    add_page_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    document.add_page_break()


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        values = [item.strip() for item in line.strip().strip("|").split("|")]
        rows.append(values)
    return rows


def add_table(document: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    # The second markdown row contains alignment markers.
    data = [rows[0], *rows[2:]]
    width = max(len(row) for row in data)
    table = document.add_table(rows=0, cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for row_index, values in enumerate(data):
        cells = table.add_row().cells
        for column_index in range(width):
            value = values[column_index] if column_index < len(values) else ""
            paragraph = cells[column_index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value)
            set_cell_margins(cells[column_index])
            cells[column_index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in paragraph.runs:
                run.font.size = Pt(8)
        if row_index == 0:
            set_repeat_table_header(table.rows[-1])
            for cell in cells:
                shade(cell, GREEN)
                set_cell_text_color(cell, WHITE)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
        elif row_index % 2 == 0:
            for cell in cells:
                shade(cell, MIST)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def render_markdown(document: Document, source: str) -> None:
    lines = source.splitlines()
    index = 0
    # The cover already contains the document title and metadata. Begin at section 1.
    while index < len(lines) and not lines[index].startswith("## 1."):
        index += 1

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped or stripped == "---":
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[index + 1]):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, parse_table(table_lines))
            continue

        heading = re.match(r"^(#{2,4})\s+(.*)$", stripped)
        if heading:
            hashes, title = heading.groups()
            level = len(hashes) - 1
            if level == 1 and not title.startswith("1."):
                document.add_page_break()
            paragraph = document.add_heading(level=level)
            add_inline(paragraph, title)
            index += 1
            continue

        checkbox = re.match(r"^- \[([ xX])\]\s+(.*)$", stripped)
        if checkbox:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, ("☒ " if checkbox.group(1).lower() == "x" else "☐ ") + checkbox.group(2))
            index += 1
            continue

        bullet = re.match(r"^-\s+(.*)$", stripped)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, bullet.group(1))
            index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, numbered.group(1))
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate == "---"
                or candidate.startswith("#")
                or candidate.startswith("|")
                or re.match(r"^-\s+", candidate)
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines))


def main() -> None:
    document = Document()
    configure_styles(document)
    configure_page(document.sections[0])
    document.sections[0].different_first_page_header_footer = True
    add_cover(document)

    body_section = document.add_section(WD_SECTION.CONTINUOUS)
    configure_page(body_section)
    add_header_footer(body_section)
    add_toc(document)
    render_markdown(document, SOURCE.read_text(encoding="utf-8"))

    properties = document.core_properties
    properties.title = "Malabar Coast Complete Product Delivery, Operations and Handover Report"
    properties.subject = "Client product delivery and operational handover"
    properties.author = "Codrant Labs"
    properties.keywords = "Malabar Coast, product delivery, handover, administration, CMS, Stripe, Supabase, Sanity"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"Created {OUTPUT} ({OUTPUT.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
